"""CSV -> PyCHAM J.xlsx (sheet "J"). Run with the PyCHAM venv python (has openpyxl).

Usage: python -I write_j_xlsx.py J.csv J.xlsx

The CSV has row 1 = times [s] and one row per MCM J index 0..61. After writing, the
file is read back the way PyCHAM's J_value_file_open.py reads it and compared.
"""
import csv
import sys

import openpyxl


def main(src, dst):
    with open(src, newline="") as fh:
        rows = [[float(v) for v in row] for row in csv.reader(fh)]
    if len(rows) != 63 or len({len(r) for r in rows}) != 1:
        raise SystemExit(f"expected 63 equal-length rows (times + J0..J61), got {len(rows)}")
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "J"
    for row in rows:
        sheet.append(row)
    wb.save(dst)
    back = list(openpyxl.load_workbook(filename=dst)["J"].iter_rows(values_only=True))
    # openpyxl stores ~16 significant digits: allow a relative 1e-14
    if len(back) != len(rows) or any(
            len(b) != len(a) or any(abs(float(y) - x) > 1e-14 * abs(x) for x, y in zip(a, b))
            for a, b in zip(rows, back)):
        raise SystemExit("J.xlsx read-back differs from the CSV")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
